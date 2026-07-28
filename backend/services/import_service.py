import csv
import io

from openpyxl import load_workbook
from pydantic import BaseModel, EmailStr, Field, ValidationError


class ImportRow(BaseModel):
    npm: str = Field(..., pattern=r"^\d{6,20}$")
    email: EmailStr


class ParsedImport(BaseModel):
    rows: list[ImportRow]
    invalid_rows: list[dict]
    total_rows: int


def parse_import_file(filename: str, content: bytes) -> ParsedImport:
    """Normalize a .csv or .xlsx student import file into validated rows.

    Both formats require `npm` and `email` header columns. Invalid rows are
    collected (1-based data row number + reason) instead of failing the batch;
    duplicate npm/email within the file keep the first occurrence.
    """
    if filename.lower().endswith(".csv"):
        raw_rows = _rows_from_csv(content)
    else:
        raw_rows = _rows_from_xlsx(content)

    rows: list[ImportRow] = []
    invalid_rows: list[dict] = []
    seen_npm: set[str] = set()
    seen_email: set[str] = set()

    for idx, raw in enumerate(raw_rows, start=1):
        npm = str(raw.get("npm") or "").strip()
        email = str(raw.get("email") or "").strip()
        try:
            row = ImportRow(npm=npm, email=email)
        except ValidationError as exc:
            reasons = "; ".join(f"{e['loc'][0]}: {e['msg']}" for e in exc.errors())
            invalid_rows.append({"row": idx, "reason": reasons})
            continue

        if row.npm in seen_npm:
            invalid_rows.append({"row": idx, "reason": "duplicate npm within file"})
            continue
        if row.email.lower() in seen_email:
            invalid_rows.append({"row": idx, "reason": "duplicate email within file"})
            continue

        seen_npm.add(row.npm)
        seen_email.add(row.email.lower())
        rows.append(row)

    return ParsedImport(rows=rows, invalid_rows=invalid_rows, total_rows=len(raw_rows))


def _rows_from_csv(content: bytes) -> list[dict]:
    decoded = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(decoded))
    if reader.fieldnames is None:
        raise ValueError("File is empty")
    headers = {h.strip().lower() for h in reader.fieldnames if h}
    _require_headers(headers)
    return [
        { (k or "").strip().lower(): v for k, v in row.items() }
        for row in reader
    ]


def _rows_from_xlsx(content: bytes) -> list[dict]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    if sheet is None:
        raise ValueError("Workbook has no active sheet")

    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        raise ValueError("File is empty")

    headers = [str(h).strip().lower() if h is not None else "" for h in header_row]
    _require_headers(set(headers))

    rows = []
    for row in rows_iter:
        if row is None or all(cell is None for cell in row):
            continue
        rows.append({headers[i]: row[i] for i in range(min(len(headers), len(row))) if headers[i]})
    workbook.close()
    return rows


def _require_headers(headers: set[str]) -> None:
    missing = {"npm", "email"} - headers
    if missing:
        raise ValueError(f"Missing required column(s): {', '.join(sorted(missing))}")